"""API-owned private temporary files; only metadata is stored in SQLite."""
import csv
import io
import json
import math
import os
import sqlite3
import time
import uuid
import zipfile
from pathlib import Path
from copy import deepcopy
from contextlib import contextmanager

MAX_BYTES = 8_000_000
MAX_ROWS = 10_000


class UploadError(ValueError):
    pass


def parse_upload(data, filename, schema):
    if not data or len(data) > MAX_BYTES: raise UploadError('File must contain 1 to 8,000,000 bytes')
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == '.csv':
            csv.field_size_limit(MAX_BYTES)
            rows = csv.reader(io.StringIO(data.decode('utf-8-sig')), strict=True)
            return _parse_rows(rows, schema)
        if suffix != '.xlsx': raise UploadError('Only .csv and .xlsx files are supported')
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 40_000_000:
                raise UploadError('Excel expanded size exceeds limit')
            if any('vbaproject' in e.filename.lower() or 'externallinks/' in e.filename.lower() for e in entries):
                raise UploadError('Macros and external links are not supported')
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) != 1: raise UploadError('Excel must contain exactly one worksheet')
            sheet = book.worksheets[0]
            if sheet.max_row and sheet.max_row > MAX_ROWS + 1: raise UploadError('Too many rows')
            if sheet.max_column and sheet.max_column > len(schema['fields']): raise UploadError('Unknown columns')
            def cells():
                for row in sheet.iter_rows():
                    if any(c.data_type == 'f' for c in row): raise UploadError('Excel formulas are not supported')
                    yield [c.value for c in row]
            return _parse_rows(cells(), schema)
        finally: book.close()
    except UploadError: raise
    except Exception:
        raise UploadError('Unable to parse file; use UTF-8 CSV or a valid XLSX workbook') from None


def _parse_rows(rows, schema):
    rows = iter(rows)
    header = next(rows, [])
    if not header or any(not isinstance(x, str) or not x.strip() for x in header): raise UploadError('Missing or invalid column names')
    header = [x.strip() for x in header]
    fields = {f['name']: f for f in schema['fields']}
    if len(header) != len(set(header)): raise UploadError('Duplicate columns')
    if set(header) - fields.keys(): raise UploadError('Unknown columns')
    if any(f['required'] and name not in header for name, f in fields.items()): raise UploadError('Missing required columns')
    result = []
    for number, cells in enumerate(rows, 2):
        if number > MAX_ROWS + 1: raise UploadError('Too many rows (maximum 10,000)')
        if not cells or all(c is None or c == '' for c in cells): continue
        if len(cells) != len(header): raise UploadError(f'Row {number}: column count mismatch')
        row = {}
        for name, raw in zip(header, cells):
            f = fields[name]
            try:
                if raw is None or raw == '':
                    if f['required']: raise ValueError()
                    continue
                kind = f['type']
                if kind == 'string':
                    if not isinstance(raw, str): raise ValueError()
                    value = raw
                elif kind == 'boolean':
                    if type(raw) is bool: value = raw
                    elif isinstance(raw, str) and raw.strip().lower() in ('true', 'false'): value = raw.strip().lower() == 'true'
                    else: raise ValueError()
                elif kind == 'number':
                    if type(raw) is bool: raise ValueError()
                    value = float(raw)
                    if not math.isfinite(value): raise ValueError()
                elif kind == 'json':
                    value = json.loads(raw, parse_constant=lambda _: (_ for _ in ()).throw(ValueError()))
                else: raise ValueError()
                if name == 'messages':
                    if not isinstance(value, list) or not value or not all(isinstance(m, dict) and m.get('role') in ('system','user','assistant') and isinstance(m.get('content'), str) for m in value): raise ValueError()
                row[name] = value
            except Exception:
                raise UploadError(f'Row {number}, column {name}: expected {f["type"]}' + (' with valid conversation messages' if name == 'messages' else '')) from None
        result.append(row)
    if not result: raise UploadError('Dataset must contain at least one sample')
    return result


class UploadStore:
    def __init__(self, root, ttl=3600):
        self.root, self.ttl = Path(root), ttl
        self.root.mkdir(parents=True, exist_ok=True, mode=0o700)
        with self.connect() as db:
            db.execute('CREATE TABLE IF NOT EXISTS uploads (id TEXT PRIMARY KEY, project TEXT, dataset TEXT, version INTEGER, count INTEGER, expires REAL, job TEXT, experiment TEXT, deleted INTEGER DEFAULT 0)')
        os.chmod(self.root / 'metadata.sqlite', 0o600)

    @contextmanager
    def connect(self):
        db = sqlite3.connect(self.root / 'metadata.sqlite', timeout=30)
        db.row_factory = sqlite3.Row
        try:
            with db: yield db
        finally: db.close()

    def create(self, project, dataset, version, schema, rows):
        ident = 'upload-' + uuid.uuid4().hex
        path = self.root / (ident + '.json')
        value = {**deepcopy(schema), 'rows': rows}
        fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        try:
            with os.fdopen(fd, 'w') as f: json.dump(value, f, allow_nan=False)
            expires = time.time() + self.ttl
            with self.connect() as db:
                db.execute('INSERT INTO uploads (id,project,dataset,version,count,expires) VALUES (?,?,?,?,?,?)', (ident,project,dataset,version,len(rows),expires))
        except BaseException:
            path.unlink(missing_ok=True)
            raise
        return {'id':ident, 'row_count':len(rows), 'dataset_id':dataset, 'dataset_version':version, 'expires_at':expires}

    def _row(self, db, project, ident):
        row = db.execute('SELECT * FROM uploads WHERE id=? AND project=?', (ident, project)).fetchone()
        if row is None or row['deleted'] or row['expires'] <= time.time(): raise UploadError('Upload not found or expired')
        return row

    def claim(self, project, ident, dataset, version, job, experiment):
        with self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            row = self._row(db,project,ident)
            if (row['dataset'], row['version']) != (dataset,version): raise UploadError('Upload schema version mismatch')
            if row['job'] and (row['job'],row['experiment']) != (job,experiment): raise UploadError('Upload already bound to another run')
            db.execute('UPDATE uploads SET job=?, experiment=?, expires=? WHERE id=?', (job,experiment,time.time()+86400,ident))
            return row['count']

    def read(self, project, ident, experiment):
        with self.connect() as db:
            row = self._row(db,project,ident)
            if row['experiment'] != experiment: raise UploadError('Upload is not bound to this run')
            try: return json.loads((self.root / (row['id'] + '.json')).read_text())
            except FileNotFoundError: raise UploadError('Upload content has been deleted') from None

    def delete(self, project, ident):
        with self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            row=db.execute('SELECT * FROM uploads WHERE id=? AND project=?',(ident,project)).fetchone()
            if row:
                (self.root / (row['id']+'.json')).unlink(missing_ok=True)
                db.execute('UPDATE uploads SET deleted=1 WHERE id=?',(ident,))

    def reconcile(self, terminal):
        with self.connect() as db: rows=db.execute('SELECT * FROM uploads WHERE deleted=0').fetchall()
        for row in rows:
            state = terminal(row['job']) if row['job'] else False
            orphan = row['job'] and state is None and time.time() > row['expires'] - 86400 + 60
            if row['expires'] <= time.time() or state is True or orphan:
                self.delete(row['project'],row['id'])
        # Crash between file creation and metadata commit: remove old orphan files.
        with self.connect() as db: known={r[0] for r in db.execute('SELECT id FROM uploads')}
        for path in self.root.glob('upload-*.json'):
            if path.stem not in known and time.time()-path.stat().st_mtime > self.ttl: path.unlink(missing_ok=True)


def get_upload_store():
    return UploadStore(os.getenv('EVAL_UPLOAD_DIR','data/uploads'), int(os.getenv('EVAL_UPLOAD_TTL_SECONDS','3600')))
